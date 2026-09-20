"""Builds docs/Writeup.pptx — the 4-slide executive write-up (deliverable 4.2).

    python docs/make_deck.py

Design brief: leadership audience, low word count, one idea per band, real numbers only.
Every figure here is traceable to the code or to docs/EVAL.md; nothing is aspirational unless
it sits under a dashed "not built" card. A fit check (Arial metrics via PIL) fails the build
if any text would overflow its shape or stray outside the slide margins.
"""
import sys
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as C
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches as I, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "Writeup.pptx"
DIAGRAM = HERE / "diagram_flow.png"
FONT = "Arial"
W, H = 13.333, 7.5
MARGIN = 0.55

INK = C(0x0F, 0x17, 0x2A)
MUTE = C(0x64, 0x74, 0x8B)
FAINT = C(0x94, 0xA3, 0xB8)
WHITE = C(0xFF, 0xFF, 0xFF)
RULE = C(0xE2, 0xE8, 0xF0)
WASH = C(0xF6, 0xF8, 0xFB)
BLUE = C(0x1D, 0x4E, 0xD8)
BLUE_BG = C(0xEE, 0xF3, 0xFE)
RUST = C(0xC2, 0x41, 0x0C)
RUST_BG = C(0xFF, 0xF3, 0xEA)
GREEN = C(0x04, 0x78, 0x57)
GREEN_BG = C(0xE6, 0xF5, 0xEF)
RED = C(0xB9, 0x1C, 0x1C)
RED_BG = C(0xFD, 0xEC, 0xEC)
PAPER = C(0xE2, 0xE8, 0xF0)

prs = Presentation()
prs.slide_width, prs.slide_height = I(W), I(H)
TOTAL = 4


# ----------------------------------------------------------------- primitives
def _para(tf, i, d):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = d.get("align", PP_ALIGN.LEFT)
    p.space_after = Pt(d.get("after", 3))
    p.line_spacing = d.get("lead", 1.0)
    for txt, bold, color in d.get("runs") or [(d["t"], d.get("bold", False), d.get("color", INK))]:
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.bold, r.font.name = Pt(d.get("size", 11)), bold, FONT
        r.font.color.rgb = color
    if d.get("bullet"):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(I(0.16))))
        pPr.set("indent", str(-int(I(0.16))))
        etree.SubElement(pPr, qn("a:buChar")).set("char", "–")
    return p


def fill(tf, paras, anchor=MSO_ANCHOR.TOP, pad=(0.16, 0.12, 0.14, 0.1)):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (I(m) for m in pad)
    for i, d in enumerate(paras):
        _para(tf, i, d)


def card(slide, x, y, w, h, paras, bg=WASH, edge=None, dashed=False, anchor=MSO_ANCHOR.TOP,
         pad=(0.16, 0.12, 0.14, 0.1), radius=0.045):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = bg
    if edge:
        s.line.color.rgb = edge
        s.line.width = Pt(1)
        if dashed:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    fill(s.text_frame, paras, anchor, pad)
    return s


def label(slide, x, y, w, t, color=FAINT, size=9):
    b = slide.shapes.add_textbox(I(x), I(y), I(w), I(0.22))
    fill(b.text_frame, [{"t": t.upper(), "size": size, "bold": True, "color": color, "after": 0}], pad=(0, 0, 0, 0))
    return b


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    b = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    fill(b.text_frame, paras, anchor, pad=(0, 0, 0, 0))
    return b


def chip(slide, x, y, w, t, bg=BLUE_BG, fg=BLUE, h=0.28, size=9.5, dashed=False, edge=None):
    return card(slide, x, y, w, h, [{"t": t, "size": size, "bold": True, "color": fg, "after": 0,
                                     "align": PP_ALIGN.CENTER}],
                bg=bg, edge=edge, dashed=dashed, anchor=MSO_ANCHOR.MIDDLE, pad=(0.05, 0, 0.05, 0), radius=0.5)


def stat(slide, x, y, w, h, big, cap, color=INK, bg=WASH, bigsize=30):
    return card(slide, x, y, w, h,
                [{"t": big, "size": bigsize, "bold": True, "color": color, "after": 2, "align": PP_ALIGN.CENTER},
                 {"t": cap, "size": 9.5, "color": MUTE, "after": 0, "align": PP_ALIGN.CENTER, "lead": 0.95}],
                bg=bg, anchor=MSO_ANCHOR.MIDDLE, pad=(0.08, 0.06, 0.08, 0.06))


def banner(slide, x, y, w, h, runs, bg=INK, size=11.5):
    return card(slide, x, y, w, h, [{"runs": runs, "size": size, "after": 0, "lead": 1.15}],
                bg=bg, anchor=MSO_ANCHOR.MIDDLE, pad=(0.22, 0.08, 0.2, 0.08), radius=0.09)


def frame(title, kicker, n, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, MARGIN, 0.34, W - 2 * MARGIN, 0.5, [{"t": title, "size": 26, "bold": True, "after": 0}])
    text(s, MARGIN, 0.92, W - 2 * MARGIN, 0.28, [{"t": kicker, "size": 11.5, "color": MUTE, "after": 0}])
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(MARGIN), I(1.26), I(W - 2 * MARGIN), I(0.012))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RULE
    ln.line.fill.background()
    ln.shadow.inherit = False
    text(s, MARGIN, 7.12, 8.0, 0.24, [{"t": "Junior AI Investigator  ·  AI-first case review for fraud investigators",
                                       "size": 8.5, "color": FAINT, "after": 0}])
    text(s, W - MARGIN - 1.0, 7.12, 1.0, 0.24, [{"t": f"{n} / {TOTAL}", "size": 8.5, "color": FAINT, "after": 0,
                                                 "align": PP_ALIGN.RIGHT}])
    s.notes_slide.notes_text_frame.text = notes
    return s


def bl(t, size=11, after=5, runs=None):
    d = {"t": t, "size": size, "bullet": True, "after": after, "lead": 1.05}
    if runs:
        d["runs"] = runs
    return d


# ================================================================== slide 1
s = frame("Investigators should open to a queue that is already triaged",
          "The user · the job to be done · what I built and what I left out", 1,
          "THE USER is a fraud investigator and their supervisor. Every morning an engine hands them a queue of referrals, "
          "and most of those turn out benign. Today they paste a claim number into a legacy tool and read nine signals a case, "
          "one case at a time. The cost is not the fraud they find, it is the hours spent on the cases that were never fraud.\n\n"
          "THE JOB TO BE DONE: decide quickly, and defensibly, which referrals deserve human time; and on the ones that do, "
          "get the evidence, the reasoning and a recommended next step without assembling it by hand.\n\n"
          "THE DATA SHAPED THE BUILD. The 50-case sample splits three ways: 8 cases extreme on every signal, 9 genuinely "
          "ambiguous, 33 clean. That split is the whole product thesis: deterministic rules are enough for the 8 and the 33, "
          "so the expensive reasoning belongs to the 9. Two patterns no row-by-row engine can see: three cases trip no flag at "
          "all yet carry five elevated signals, and one claim number (LTC-2034786) sits on both a clean case and a critical one.\n\n"
          "WHAT I CUT AND WHY: auth, multi-tenancy and deployment are explicitly out of scope in the brief. I did not train a "
          "model because the data has no fraud labels, so any accuracy number would be invented. A provider/member graph needs "
          "IDs the CSV does not contain; the code supports them the day a dataset does.")

card(s, MARGIN, 1.5, 6.05, 1.42,
     [{"t": "The user", "size": 10, "bold": True, "color": BLUE, "after": 6},
      {"t": "A fraud investigator and their supervisor, facing a daily queue of engine referrals that mostly turn out benign.",
       "size": 12.5, "after": 0, "lead": 1.15}],
     bg=BLUE_BG)
card(s, MARGIN + 6.17, 1.5, 6.06, 1.42,
     [{"t": "The job to be done", "size": 10, "bold": True, "color": BLUE, "after": 6},
      {"t": "Decide fast, and defensibly, which referrals deserve human time — and on those, get the evidence and the next step without assembling it by hand.",
       "size": 12.5, "after": 0, "lead": 1.15}],
     bg=BLUE_BG)

label(s, MARGIN, 3.14, 6.0, "What the 50-case sample showed")
stat(s, MARGIN, 3.42, 3.97, 1.08, "8", "extreme on every signal", color=RED, bg=RED_BG)
stat(s, MARGIN + 4.09, 3.42, 3.97, 1.08, "9", "genuinely ambiguous", color=RUST, bg=RUST_BG)
stat(s, MARGIN + 8.18, 3.42, 4.05, 1.08, "33", "clean", color=GREEN, bg=GREEN_BG)
banner(s, MARGIN, 4.64, 12.23, 0.62,
       [("So: rules take the obvious 41. Agents and humans spend their time on the 9 — plus the links no single row reveals: ",
         True, WHITE),
        ("3 cases trip no flag yet carry five elevated signals, and one claim number sits on a clean case and a critical one.",
         False, C(0xCB, 0xD5, 0xE1))], size=11)

label(s, MARGIN, 5.5, 6.0, "Built in the prototype", color=BLUE)
built = ["Morning briefing", "Per-case assessment", "Evidence + agent trace", "Accept / reject / note",
         "Bulk clear, guarded", "Supervisor escalation", "Rules + backtest", "Blocklist",
         "Outcome feedback", "Audit log"]
for i, t in enumerate(built):
    chip(s, MARGIN + (i % 5) * 1.63, 5.78 + (i // 5) * 0.37, 1.53, t, size=8.5)

label(s, 8.85, 5.5, 3.93, "Left out on purpose", color=MUTE)
for i, (t, why) in enumerate([("Auth, multi-tenancy, deploy", "out of scope in the brief"),
                              ("A trained ML model", "no labels — accuracy would be invented"),
                              ("Provider / member graph", "no IDs in the CSV; code is ready")]):
    card(s, 8.85, 5.78 + i * 0.39, 3.93, 0.34,
         [{"runs": [(t + "  ", True, MUTE), (why, False, FAINT)], "size": 9, "after": 0}],
         bg=WHITE, edge=RULE, anchor=MSO_ANCHOR.MIDDLE, pad=(0.12, 0, 0.1, 0), radius=0.12)

# ================================================================== slide 2
s = frame("Rules set the depth; agents reason only where it pays",
          "How the system reasons about a case, stays grounded, and scales", 2,
          "NOT ONE LLM CALL, AND NOT AGENTS EVERYWHERE. Deterministic rules score every case first — 17 rules whose thresholds "
          "come from natural breaks in the data, not from guesses — and a router with no LLM in it reads the fired rules and "
          "picks the depth. 58% of the queue fired nothing and is unlinked: those share one cheap batched call, five cases at a "
          "time. 16% are rule-CRITICAL, where the facts are not in doubt, so buying a debate would be waste: one verifier call. "
          "The remaining 26% are the ambiguous or linked cases, and only those get the full crew: specialists in parallel for the "
          "domains that actually fired (billing, collusion and linkage, geography and utilisation, one tool each), an adversarial "
          "challenger that argues the benign reading, a verifier on the strong model tier, and a three-persona panel that convenes "
          "only when the case is still contested.\n\n"
          "TOOL USE, NOT RETRIEVAL. Nine read-only tools query the real record — peer statistics, similar cases, linked claim "
          "numbers, notes, the rule catalogue, cluster context. No vector store and no RAG: the corpus is 50 structured rows, so "
          "SQL answers exactly, and embeddings would only add a way to be wrong.\n\n"
          "GROUNDING IS CODE, NOT PROMPTING. Every field=value the model cites is re-checked against the record before it is "
          "stored; the model may move the score by at most 20 points and the level is derived from the number, not chosen; it "
          "cannot recommend clearing a HIGH/CRITICAL case or a confirmed-fraud lookalike; and when the model fails, the case says "
          "LLM unavailable instead of showing invented text. Measured across the 50 cases: 10/10 behavioural invariants, 5/5 "
          "counterfactual and prompt-injection tests, 277/277 stored indicators re-verify.\n\n"
          "SCALE. The router is free and O(1), so spend lands on the hard tail — and a real queue has a larger clean share than "
          "this sample. Against a single-judge baseline the crew agrees on the band 100% of the time with 27% fewer tokens, and a "
          "new note re-runs two calls instead of five because specialists are reused by evidence hash. The honest limits are "
          "SQLite, in-process workers and O(n) similarity.")

pic = s.shapes.add_picture(str(DIAGRAM), I(MARGIN), I(1.46), height=I(5.42))
pic.line.color.rgb, pic.line.width = RULE, Pt(0.75)
DX = MARGIN + pic.width / 914400 + 0.42
DW = W - MARGIN - DX

card(s, DX, 1.46, DW, 1.52,
     [{"t": "How it reasons", "size": 10, "bold": True, "color": RUST, "after": 6},
      {"runs": [("Not a single call, and not agents everywhere. ", True, INK),
                ("Deterministic rules score and route every case; depth is bought only where the evidence is genuinely unclear.",
                 False, INK)], "size": 11, "after": 6, "lead": 1.12},
      {"runs": [("58% ", True, GREEN), ("one batched call  ·  ", False, MUTE),
                ("16% ", True, RED), ("verifier only  ·  ", False, MUTE),
                ("26% ", True, RUST), ("specialists → challenger → verifier → panel if contested", False, MUTE)],
       "size": 10, "after": 0}],
     bg=RUST_BG)

card(s, DX, 3.06, DW, 1.04,
     [{"t": "Tool use, not retrieval", "size": 10, "bold": True, "color": BLUE, "after": 6},
      {"runs": [("9 read-only tools query the live record", True, INK),
                (" — peers, similar cases, linked claims, notes, rules. No vector store, no RAG: 50 structured rows answer exactly in SQL.",
                 False, INK)], "size": 11, "after": 0, "lead": 1.12}],
     bg=BLUE_BG)

card(s, DX, 4.18, DW, 1.78,
     [{"t": "Grounded by code, not by prompt", "size": 10, "bold": True, "color": WHITE, "after": 6},
      bl("", 10.5, 4, runs=[("Every cited field = value re-verified against the record", False, PAPER)]),
      bl("", 10.5, 4, runs=[("Score moves ≤ ±20; the risk level is derived, not chosen", False, PAPER)]),
      bl("", 10.5, 4, runs=[("Cannot clear a HIGH/CRITICAL case or a confirmed-fraud lookalike", False, PAPER)]),
      bl("", 10.5, 6, runs=[("On failure it says “LLM unavailable” — never invented text", False, PAPER)]),
      {"runs": [("10/10", True, C(0x6E, 0xE7, 0xB7)), (" invariants   ", False, C(0x94, 0xA3, 0xB8)),
                ("5/5", True, C(0x6E, 0xE7, 0xB7)), (" adversarial tests   ", False, C(0x94, 0xA3, 0xB8)),
                ("277/277", True, C(0x6E, 0xE7, 0xB7)), (" indicators re-verified", False, C(0x94, 0xA3, 0xB8))],
       "size": 10, "after": 0}],
     bg=INK)

card(s, DX, 6.08, DW, 0.9,
     [{"t": "Beyond 50 cases", "size": 10, "bold": True, "color": MUTE, "after": 5},
      {"runs": [("Router is free and O(1) — spend lands on the hard tail. ", False, INK),
                ("100% band agreement with a single judge at 27% fewer tokens. ",
                 True, INK),
                ("Limits: SQLite, in-process workers, O(n) similarity.", False, MUTE)], "size": 10, "after": 0,
       "lead": 1.12}],
     bg=WHITE, edge=RULE)

# ================================================================== slide 3
s = frame("What I cut, what I would do next, and the risks I can name",
          "Trade-offs · deliberate omissions · where this could go wrong", 3,
          "THE HONEST HEADLINE: the data has no fraud labels, so nothing in this build claims accuracy. What I can evidence is "
          "behaviour — invariants that hold whatever the model says, grounding that survives an independent re-check, and "
          "adversarial tests that pass. Everything else on this slide is named risk, not solved risk.\n\n"
          "FALSE NEGATIVES ARE THE ASYMMETRIC RISK. A false positive costs an investigator an hour; a false negative pays a "
          "fraudulent claim and never shows up in a metric. The guardrails stop the AI clearing anything risky, but they cannot "
          "find what the rules never flagged — which is exactly why the safety net is the first item on the roadmap rather than "
          "a nice-to-have.\n\n"
          "COST AND LATENCY, MEASURED NOT GUESSED. The crew uses 27% fewer tokens than a single judge but costs more in "
          "reference dollars ($0.44 vs $0.12 across 50 cases) because the verifier runs on the strong tier — I would rather state "
          "that plainly than claim the architecture is cheaper on every axis. At roughly 20 seconds a call on free models, this is "
          "a morning batch job, not a per-click interaction, and the UI is built around that.\n\n"
          "ABUSE AND FAIRNESS. Notes are untrusted input and injection tests pass. Bulk clear previews what it cannot touch. A "
          "blocklist entry covering more than a quarter of the queue demands explicit confirmation. On fairness, distance and "
          "weekend-billing signals can proxy for rural or shift-work providers — worth auditing before any real deployment.\n\n"
          "MY OWN EVALS FOUND WEAKNESSES: the challenger rates nearly every case 'medium' plausibility, and confidence barely "
          "separates contested from uncontested cases. Both are reported rather than hidden.\n\n"
          "DEFERRED ON PURPOSE: I prototyped a statistical anomaly detector and shelved it — at 50 cases its flagged set changed "
          "with the random seed, and a detector that unstable erodes trust faster than it finds fraud.")

label(s, MARGIN, 1.46, 6.0, "Risks I can name", color=RED)
risks = [("Accuracy", "No labels in the data — so no accuracy claim. Backtests measure agreement, not fraud.", RED),
         ("False negatives", "The asymmetric risk: only discovered misses can ever be measured.", RED),
         ("Cost", "27% fewer tokens than one judge, but 3.6× the reference dollars — the verifier runs strong-tier.", RUST),
         ("Latency", "~20s a call on free models. A morning batch, not a per-click interaction.", RUST),
         ("Trust", "Cited fields are verified; the prose around them is not. It can still read wrong.", RUST),
         ("Abuse", "Notes are untrusted input; bulk clear and broad blocklist entries are gated.", MUTE),
         ("Fairness", "Distance and weekend signals can proxy for rural or shift-work providers.", MUTE)]
for i, (k, v, col) in enumerate(risks):
    card(s, MARGIN, 1.74 + i * 0.72, 6.05, 0.64,
         [{"runs": [(k + "   ", True, col), (v, False, INK)], "size": 10.5, "after": 0, "lead": 1.12}],
         bg=WHITE, edge=RULE, anchor=MSO_ANCHOR.MIDDLE, pad=(0.16, 0, 0.14, 0))

label(s, 6.98, 1.46, 5.8, "Deliberately cut, and why", color=MUTE)
card(s, 6.98, 1.74, 5.8, 1.74,
     [bl("Auth, multi-tenancy, deployment — out of scope in the brief", 10.5, 6),
      bl("A trained model — no labels, so any accuracy figure would be fiction", 10.5, 6),
      bl("Provider / member graph — the CSV has no such IDs; the code is ready for them", 10.5, 6),
      bl("External enrichment — compelling, but not the core of the AI experience", 10.5, 0)],
     bg=WASH)

label(s, 6.98, 3.62, 5.8, "With more time, in priority order", color=BLUE)
nxt = [("1", "A false-negative safety net", "miss-rate from confirmed outcomes · look-back that reopens cleared lookalikes · an independent adversarial reviewer"),
       ("2", "Learn from outcomes", "re-fit weights and thresholds on confirmed fraud instead of fixed calibration"),
       ("3", "A labelled backtest", "the moment real outcomes exist, replace synthetic-label evaluation")]
yy = 3.9
for num, head, body in nxt:
    h = 0.8 if num == "1" else 0.64
    card(s, 6.98, yy, 5.8, h,
         [{"runs": [(num + "   ", True, BLUE), (head + "  ", True, INK), (body, False, MUTE)], "size": 10.5,
           "after": 0, "lead": 1.12}],
         bg=BLUE_BG, anchor=MSO_ANCHOR.MIDDLE)
    yy += h + 0.1

card(s, 6.98, yy + 0.02, 5.8, 0.66,
     [{"runs": [("Prototyped, then deferred.  ", True, MUTE),
                ("A statistical anomaly detector for the silent cases: at 50 cases its flagged set moved with the random seed, so I would not ship it.",
                 False, MUTE)], "size": 10, "after": 0, "lead": 1.12}],
     bg=WHITE, edge=MUTE, dashed=True, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================== slide 4
s = frame("The AI advises. Only a human decides.",
          "How investigators trust, validate and override the system", 4,
          "THE HARD LINE: no agent in this system can change the state of a case. The AI ranks, explains, recommends and drafts; "
          "clearing, escalating, declining, confirming fraud and deleting a note are human actions, each one requiring a written "
          "reason and each one landing in an audit log. That is enforced in the service layer, not requested in a prompt.\n\n"
          "TRUST IS BUILT FROM FOUR THINGS. First, both scores are shown side by side, so the investigator always sees where the "
          "model moved away from the deterministic rules and by how much. Second, every indicator cites the field and value it "
          "came from, and can be accepted or rejected with one click — a rejected finding is not re-asserted on the next run. "
          "Third, the full agent trace is visible: which specialists ran, what the challenger argued, what the guardrails removed "
          "and why, and what each step cost. Fourth, every action carries a reason into the audit log.\n\n"
          "NOTES ARE THE STEERING WHEEL. An investigator's note changes the next assessment — the case is marked stale the moment "
          "its inputs change — and the chatbot can quote it. Deleting a note purges it everywhere: assessments, chat history, "
          "related-case context. Notes are also treated as untrusted text, so an instruction hidden inside one is not obeyed; "
          "that is covered by passing injection tests.\n\n"
          "AUTHORITY IS SPLIT. Investigators propose blocklist entries; supervisors enact them. Only a supervisor can override an "
          "auto-decline, and only with a written reason that becomes a case note and counts against the entry that fired it as a "
          "false-positive signal. Escalations carry an AI-drafted handoff, and the supervisor approves or returns them.\n\n"
          "DEMO PATH: add a note and watch the assessment go stale, ask the chatbot about it, then delete it; run a guarded bulk "
          "clear; decline a case from the blocklist and override it as the supervisor.")

flow = [("AI proposes", "summary · indicators · risk · next step", RUST, RUST_BG),
        ("Code verifies", "re-ground · clamp ±20 · refuse", INK, WHITE),
        ("Human decides", "clear · escalate · note · override", GREEN, GREEN_BG)]
fw = 3.93
for i, (t, sub, col, bg) in enumerate(flow):
    x = MARGIN + i * (fw + 0.22)
    card(s, x, 1.46, fw, 0.86,
         [{"t": t, "size": 13, "bold": True, "color": col, "after": 3, "align": PP_ALIGN.CENTER},
          {"t": sub, "size": 9.5, "color": MUTE, "after": 0, "align": PP_ALIGN.CENTER}],
         bg=bg, edge=RULE if bg == WHITE else None, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x + fw + 0.045), I(1.78), I(0.13), I(0.22))
        a.fill.solid()
        a.fill.fore_color.rgb = FAINT
        a.line.fill.background()
        a.shadow.inherit = False

label(s, MARGIN, 2.56, 6.0, "What the AI may do — and what it may never do")
card(s, MARGIN, 2.84, 6.05, 1.7,
     [{"t": "The AI can", "size": 10, "bold": True, "color": RUST, "after": 5},
      {"t": "Rank and explain · recommend an action · draft an escalation · answer questions using read-only tools.",
       "size": 11, "after": 9, "lead": 1.12},
      {"t": "Only a human can", "size": 10, "bold": True, "color": GREEN, "after": 5},
      {"runs": [("Clear · escalate · confirm fraud · delete a note.  ", False, INK),
                ("Supervisor only: ", True, GREEN),
                ("override a decline, activate a hard rule.", False, INK)], "size": 11, "after": 0, "lead": 1.12}],
     bg=WASH)

label(s, MARGIN, 4.7, 6.0, "Notes are the steering wheel")
card(s, MARGIN, 4.98, 6.05, 1.9,
     [bl("A note changes the next assessment — the case is marked stale at once", 11, 7),
      bl("Delete it and it is purged everywhere: assessment, chat, related cases", 11, 7),
      bl("Notes are untrusted text: a hidden instruction is not obeyed (tested)", 11, 7),
      bl("Investigators propose blocklist entries; supervisors enact them", 11, 0)],
     bg=WASH)

label(s, 6.98, 2.56, 5.8, "Four things that make it checkable", color=BLUE)
trust = [("Two scores, side by side", "The deterministic rule score next to the AI score — the gap is always visible."),
         ("Citations you can reject", "Every indicator names its field and value. Accept ✓ or reject ✗; a rejected finding is not re-asserted."),
         ("The whole trace is open", "Which specialists ran, what the challenger argued, what the guardrails removed — and what it cost."),
         ("A reason on every action", "Clears, escalations and overrides all require a written reason and land in the audit log.")]
yy = 2.84
for t, b in trust:
    card(s, 6.98, yy, 5.8, 0.94,
         [{"t": t, "size": 11.5, "bold": True, "color": INK, "after": 4},
          {"t": b, "size": 10, "color": MUTE, "after": 0, "lead": 1.12}],
         bg=WHITE, edge=RULE, anchor=MSO_ANCHOR.MIDDLE)
    yy += 1.03

prs.save(OUT)


# ----------------------------------------------------------------- fit check
_CACHE = {}


def _font(size, bold):
    key = (round(size * 10), bold)
    if key not in _CACHE:
        _CACHE[key] = ImageFont.truetype(f"/System/Library/Fonts/Supplemental/Arial{' Bold' if bold else ''}.ttf", key[0])
    return _CACHE[key]


def _wrapped(txt, size, bold, width_pt):
    f, n, cur = _font(size, bold), 1, ""
    for word in txt.split(" "):
        trial = (cur + " " + word).strip()
        if f.getlength(trial) / 10 <= width_pt or not cur:
            cur = trial
        else:
            n, cur = n + 1, word
    return n


def check(path):
    bad = []
    for si, slide in enumerate(Presentation(path).slides, 1):
        for sh in slide.shapes:
            x, y, w, h = (sh.left / 914400, sh.top / 914400, sh.width / 914400, sh.height / 914400)
            if x < MARGIN - 0.01 or y < 0.3 or x + w > W - MARGIN + 0.01 or y + h > H - 0.12:
                bad.append(f"slide {si}: shape ({x:.2f},{y:.2f},{w:.2f},{h:.2f}) breaks the margins")
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            tf = sh.text_frame
            total = (tf.margin_top + tf.margin_bottom) / 12700
            avail = w * 72 - (tf.margin_left + tf.margin_right) / 12700
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                size = p.runs[0].font.size.pt
                bold = any(r.font.bold for r in p.runs)
                marl = int(p._p.pPr.get("marL", 0)) / 12700 if p._p.pPr is not None else 0
                lead = p.line_spacing or 1.0
                total += _wrapped("".join(r.text for r in p.runs), size, bold, avail - marl) * size * 1.2 * lead
                total += p.space_after.pt if p.space_after is not None else 0
            if total > h * 72 * 0.99:
                bad.append(f"slide {si}: needs {total / 72:.2f}in in a {h:.2f}in box — \"{tf.text[:46]}…\"")
    return bad


problems = check(OUT)
print("saved", OUT)
print("FIT CHECK:", "OK — all text fits, every shape inside the margins" if not problems else "FAILED")
for p in problems:
    print("   ·", p)
sys.exit(1 if problems else 0)
