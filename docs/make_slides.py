"""Builds docs/Writeup.pptx: the 3-slide write-up (deliverable 4.2).   Run:  python docs/make_slides.py
Also runs an automated fit check (Arial metrics via PIL) so text that would overflow its shape fails the build."""
import sys
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as C
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches as I, Pt

OUT = Path(__file__).resolve().parent / "Writeup.pptx"
FONT = "Arial"
INK, MUT, BRAND, SOFT, WHITE, LINE = C(0x1A, 0x1F, 0x2B), C(0x5B, 0x64, 0x78), C(0x3B, 0x4F, 0xD8), C(0xEE, 0xF0, 0xFD), C(255, 255, 255), C(0xD9, 0xDD, 0xE7)
RED, ORANGE, AMBER, GREEN = C(0xB4, 0x23, 0x18), C(0xC4, 0x52, 0x0A), C(0xA1, 0x62, 0x07), C(0x0E, 0x7A, 0x4B)
RED_BG, ORANGE_BG, AMBER_BG, GREEN_BG = C(0xFE, 0xE4, 0xE2), C(0xFF, 0xEA, 0xD5), C(0xFE, 0xF3, 0xC7), C(0xD9, 0xF5, 0xE6)

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
TOTAL = 3


def fmt_run(r, size, bold=False, color=INK, italic=False):
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, FONT
    r.font.color.rgb = color


def set_bullet(p, indent=0.15):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(I(indent))))
    pPr.set("indent", str(-int(I(indent))))
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def fill_text(tf, paras, anchor=MSO_ANCHOR.TOP, margins=(0.14, 0.1, 0.14, 0.08)):
    """paras: list of dicts {t, size, bold, color, bullet, after, align, runs:[(text,bold,color)]}"""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (I(m) for m in margins)
    for i, d in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = d.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(d.get("after", 3))
        p.line_spacing = 1.0
        for text, bold, color in d.get("runs") or [(d["t"], d.get("bold", False), d.get("color", INK))]:
            fmt_run(p.add_run(), d.get("size", 11), bold, color)
            p.runs[-1].text = text
        if d.get("bullet"):
            set_bullet(p)


def box(slide, x, y, w, h, paras, fill=SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.TOP, radius=0.05, margins=(0.14, 0.1, 0.14, 0.08), dashed=False):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
        if dashed:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    fill_text(s.text_frame, paras, anchor, margins)
    return s


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    b = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    fill_text(b.text_frame, paras, anchor, margins=(0, 0, 0, 0))
    return b


def head(t, size=13):
    return {"t": t, "size": size, "bold": True, "color": BRAND, "after": 5}


def bl(t, size=11, after=4, runs=None):
    return {"t": t, "size": size, "bullet": True, "after": after, **({"runs": runs} if runs else {})}


def slide_frame(title, sub, n, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.5, 0.3, 12.3, 0.62, [{"t": title, "size": 25, "bold": True, "after": 0}])
    text(s, 0.5, 0.93, 12.3, 0.32, [{"t": sub, "size": 13, "color": MUT, "after": 0}])
    text(s, 0.5, 7.08, 9, 0.25, [{"t": "Junior AI Investigator  ·  Take-home write-up", "size": 9, "color": MUT, "after": 0}])
    text(s, 11.8, 7.08, 1.0, 0.25, [{"t": f"{n} / {TOTAL}", "size": 9, "color": MUT, "after": 0, "align": PP_ALIGN.RIGHT}])
    s.notes_slide.notes_text_frame.text = notes
    return s


def arrow(slide, x, y, w=0.28, h=0.26):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    a.fill.solid()
    a.fill.fore_color.rgb = BRAND
    a.line.fill.background()
    return a


def tile(slide, x, y, w, h, big, label, color=BRAND, bg=SOFT):
    box(slide, x, y, w, h, [{"t": big, "size": 22, "bold": True, "color": color, "after": 0, "align": PP_ALIGN.CENTER},
                            {"t": label, "size": 10, "color": MUT, "after": 0, "align": PP_ALIGN.CENTER}], fill=bg, anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------------ slide 1
def tag(slide, x, y, w, label, built=True):
    box(slide, x, y, w, 0.24, [{"t": label, "size": 8.5, "bold": True, "color": WHITE if built else MUT, "after": 0, "align": PP_ALIGN.CENTER}],
        fill=BRAND if built else WHITE, line=None if built else MUT, dashed=not built, anchor=MSO_ANCHOR.MIDDLE, radius=0.4, margins=(0.04, 0, 0.04, 0))


s = slide_frame("Investigators open the tool to a triaged queue, not a blank one",
                "Product  ·  the user, the job to be done, and what I chose to build and to cut", 1,
                "PRODUCT. The user is a fraud investigator, and their supervisor, who starts the day with engine-referred claims that are mostly benign. "
                "The job to be done is to decide quickly and defensibly which referrals deserve human time, and for those, to get evidence, reasoning and a next step without hunting through dozens of signals.\n\n"
                "The data shaped the design. The 50 cases fall into three tiers: 8 extreme on every signal, 9 ambiguous, 33 clean. Three 'silent' cases (C1026, C1040, C1042) trip no flag but carry five elevated signals, "
                "and claim number LTC-2034786 sits on a clean case (C1001) and a critical one (C1031). So rules take the obvious cases; AI and humans spend their effort on the ambiguous middle and on cross-case links.\n\n"
                "The screenshot is the real prototype. Note '9 AI-unavailable': the free models rate-limit, and the product shows that honestly instead of inventing an assessment.\n\n"
                "DEMO PATH: Morning briefing, then a critical case, then C1001 (the linked-claim story).")
box(s, 0.5, 1.45, 5.45, 1.5, [head("The user and the job", 13),
    {"runs": [("User: ", True, INK), ("a fraud investigator, and their supervisor, facing a daily queue of engine referrals that mostly turn out benign.", False, INK)], "size": 11, "after": 5},
    {"runs": [("Job: ", True, INK), ("decide fast, and defensibly, which referrals deserve human time; on those, get evidence, reasoning and a next step without hunting through signals.", False, INK)], "size": 11, "after": 0}])
text(s, 0.5, 3.1, 5.45, 0.25, [{"t": "WHAT THE 50-CASE CSV TOLD ME", "size": 9.5, "bold": True, "color": MUT, "after": 0}])
tile(s, 0.5, 3.38, 1.72, 0.98, "8", "extreme on every signal", color=RED, bg=RED_BG)
tile(s, 2.365, 3.38, 1.72, 0.98, "9", "ambiguous middle", color=AMBER, bg=AMBER_BG)
tile(s, 4.23, 3.38, 1.72, 0.98, "33", "clean", color=GREEN, bg=GREEN_BG)
box(s, 0.5, 4.5, 5.45, 1.32, [
    bl("3 'silent' cases trip no flag, yet carry 5 elevated signals. A flag-count model calls them clean.", 10.5, 4),
    bl("Claim number LTC-2034786 sits on a clean case and a critical one. No single-row rule can see that.", 10.5, 0)],
    fill=C(0xF4, 0xF5, 0xF8), line=LINE)
box(s, 0.5, 5.95, 5.45, 1.0, [{"runs": [("Design consequence: ", True, WHITE), ("rules take the obvious; AI and humans spend their time on the ambiguous middle and the cross-case links.", False, WHITE)], "size": 11.5, "after": 0}],
    fill=INK, anchor=MSO_ANCHOR.MIDDLE)
pic = s.shapes.add_picture(str(Path(__file__).resolve().parent / "screenshot_briefing.png"), I(6.2), I(1.45), width=I(6.63))
pic.line.color.rgb, pic.line.width = LINE, Pt(1)
text(s, 6.2, 4.78, 6.63, 0.32, [{"t": "The running prototype: Morning briefing. Cases the free models could not assess show as “AI-unavailable”; nothing is invented.", "size": 9, "color": MUT, "after": 0}])
tag(s, 6.2, 5.12, 1.5, "BUILT IN PROTOTYPE")
box(s, 6.2, 5.42, 3.72, 1.53, [
    bl("Morning briefing: risk counts, exposure, clusters, abnormal firing", 10, 2),
    bl("Per-case assessment, evidence trace, accept / reject", 10, 2),
    bl("Chat with tools; notes the AI reads", 10, 2),
    bl("Bulk clear, supervisor escalation, audit log", 10, 2),
    bl("Rules tab, Blocklist, fraud-outcome feedback", 10, 0)], margins=(0.12, 0.07, 0.1, 0.05))
box(s, 10.05, 5.42, 2.78, 1.53, [
    {"t": "Cut on purpose", "size": 10.5, "bold": True, "color": MUT, "after": 3},
    bl("Auth, multi-tenancy, deploy", 10, 2),
    bl("ML training: no labels", 10, 2),
    bl("Provider / member graph: no IDs in the CSV", 10, 0)], fill=WHITE, line=MUT, dashed=True, margins=(0.12, 0.07, 0.1, 0.05))

# ------------------------------------------------------------------ slide 2
s = slide_frame("Rules set the depth; a tool-using agent crew reasons only where it pays",
                "Architecture  ·  how a case is reasoned about, kept grounded, and scaled", 2,
                "ARCHITECTURE. Not one LLM call. Deterministic rules score every case first; thresholds come from natural breaks in the data. "
                "A router with no LLM then picks a lane: 29 clean cases share one cheap batched call (5 per call), 8 critical cases get a verifier only, and 13 ambiguous or linked cases get the full crew.\n\n"
                "On the DEEP lane, specialists (billing, collusion and linkage, geography and utilization) run in parallel only where their domain fired, each with at most one tool. An adversarial challenger argues the benign case, "
                "a verifier on the strong model tier synthesises, and a three-persona panel convenes only when the case is contested. All of it runs in a small harness I wrote: budgets, tool allowlists, typed output with repair, fast-to-strong escalation, model fallback, traced spans. No agent framework, no RAG, no vector store.\n\n"
                "GROUNDING is deterministic code around the model: citations re-checked against the record, score moves capped at 20 with the level derived, no CLEAR on HIGH/CRITICAL or a confirmed-fraud lookalike, failures shown as LLM_UNAVAILABLE.\n\n"
                "MEASURED on the 50 cases (docs/EVAL.md): 10/10 invariants, 5/5 counterfactual and injection tests, 277/277 stored indicators re-verify, 100% band agreement with the single-judge baseline at 27% fewer tokens. "
                "The strong-tier verifier raises reference dollars ($0.44 vs $0.12), so I do not claim it is cheaper in dollars.\n\n"
                "SCALING: right side is explicitly not built.")
text(s, 7.6, 0.93, 5.23, 0.3, [{"runs": [("Solid = built in the prototype    ", True, BRAND), ("Dashed = production path, not built", True, MUT)], "size": 10, "after": 0, "align": PP_ALIGN.RIGHT}])
y0 = 1.42
def pipe(x, w, big, small, fill=SOFT, fg=INK, mut=MUT):
    box(s, x, y0, w, 0.8, [{"t": big, "size": 11, "bold": True, "color": fg, "after": 1, "align": PP_ALIGN.CENTER},
                           {"t": small, "size": 9, "color": mut, "after": 0, "align": PP_ALIGN.CENTER}], fill=fill, anchor=MSO_ANCHOR.MIDDLE)
pipe(0.5, 2.5, "50-case CSV", "features: log-z, travel, claim links")
arrow(s, 3.06, y0 + 0.27)
pipe(3.4, 3.0, "17 rules, R01–R17", "thresholds = natural breaks in the data")
arrow(s, 6.46, y0 + 0.27)
pipe(6.8, 2.9, "Score 0–100 + level", "top rule per domain, noisy-OR")
arrow(s, 9.76, y0 + 0.27)
pipe(10.1, 2.73, "Router, no LLM", "picks depth from the rules", fill=BRAND, fg=WHITE, mut=WHITE)
ly = 2.38
box(s, 0.5, ly, 2.75, 1.42, [{"t": "FAST  ·  29 cases", "size": 12, "bold": True, "color": GREEN, "after": 3}, {"t": "No rule fired, no link. One cheap call reviews 5 cases at once.", "size": 10.5, "after": 0}], fill=GREEN_BG)
box(s, 3.37, ly, 2.75, 1.42, [{"t": "OBVIOUS  ·  8 cases", "size": 12, "bold": True, "color": RED, "after": 3}, {"t": "Rule level CRITICAL. Verifier only: the facts are unambiguous, so no debate is bought.", "size": 10.5, "after": 0}], fill=RED_BG)
box(s, 6.24, ly, 6.59, 1.42, [{"t": "DEEP  ·  13 cases (ambiguous or linked)", "size": 12, "bold": True, "color": AMBER, "after": 3},
    {"t": "Specialists in parallel, only where their domain fired (billing · collusion + linkage · geo/utilization, ≤1 tool each)  →  adversarial challenger  →  verifier (strong tier)  →  3-persona panel only if contested.", "size": 10.5, "after": 0}], fill=AMBER_BG)
box(s, 0.5, 3.93, 12.33, 0.56, [{"runs": [("VERIFICATION BOUNDARY, deterministic code and not the model:  ", True, WHITE),
    ("citations re-checked against the record  ·  score moves ≤ 20, level derived  ·  no CLEAR on HIGH/CRITICAL or fraud lookalike  ·  failure = “LLM unavailable”", False, WHITE)], "size": 10, "after": 0}],
    fill=INK, anchor=MSO_ANCHOR.MIDDLE, radius=0.12)
ry, rh = 4.68, 2.27
text(s, 0.5, ry, 7.3, 0.22, [{"t": "MEASURED ON THE 50 CASES  (docs/EVAL.md)", "size": 9.5, "bold": True, "color": MUT, "after": 0}])
tw = 1.7
for i, (big, lab, col, bg) in enumerate([("10 / 10", "behavioural invariants", GREEN, GREEN_BG), ("5 / 5", "counterfactual tests", GREEN, GREEN_BG),
                                          ("277 / 277", "indicators re-verified", GREEN, GREEN_BG), ("−27%", "tokens vs single judge", BRAND, SOFT)]):
    tile(s, 0.5 + i * (tw + 0.1), ry + 0.28, tw, 0.86, big, lab, color=col, bg=bg)
box(s, 0.5, ry + 1.26, 7.3, 1.01, [
    bl("Same 100% band agreement as the single-judge baseline; a new note re-runs 2 calls, not 5 (specialists reused by evidence hash).", 10, 3),
    bl("Honest cost: $0.44 vs $0.12 reference, because the verifier uses the strong tier. Weak spot found: the challenger rates every case “medium”.", 10, 0)],
    fill=C(0xF4, 0xF5, 0xF8), line=LINE, margins=(0.12, 0.07, 0.1, 0.05))
tag(s, 8.0, ry, 2.2, "PRODUCTION PATH: NOT BUILT", built=False)
box(s, 8.0, ry + 0.3, 4.83, 1.97, [
    {"runs": [("Prototype limits: ", True, INK), ("SQLite, in-process workers, O(n) similarity, ~4 sequential hops (≈65 s on free models): right for a morning batch, not per click.", False, INK)], "size": 10, "after": 4},
    {"runs": [("Beyond 50 cases: ", True, INK), ("Postgres + queue workers; vector index for similarity; per-tenant budgets; streaming UI. Router stays O(1) per case, so spend lands on the hard tail (a real queue has a bigger FAST share).", False, INK)], "size": 10, "after": 0}],
    fill=WHITE, line=MUT, dashed=True, margins=(0.14, 0.09, 0.12, 0.06))

# ------------------------------------------------------------------ slide 3
s = slide_frame("The AI advises, humans decide, and the known risks are named",
                "Human in the loop  ·  trade-offs  ·  what I would do with more time", 3,
                "HUMAN IN THE LOOP. The AI never changes a case's status. Every clear, escalation, override and fraud outcome is a human action with a reason, in an audit log. "
                "Investigators see the rule score beside the AI score, check every indicator against its field and value, accept or reject it (rejections are respected next time), and read the agent trace and guardrail removals. "
                "Notes steer the next assessment, and deleting a note purges it from assessments, chat and related-case context. Supervisors approve escalations, override blocklist declines and activate hard rules; investigators can only propose them.\n\n"
                "TRADE-OFFS. No labels means no accuracy claim; backtests use synthetic labels and say so. False negatives are the costliest risk and only discovered misses can be measured. "
                "The free models rate-limit, so some cases show LLM unavailable: shown, not faked.\n\n"
                "NEXT (not built): false-negative safety net (miss-rate tracking, look-back that reopens cleared lookalikes, independent adversarial reviewer, audit sampling with a statistical bound). "
                "I prototyped an anomaly detector and deferred it: flagged sets changed with the seed at 50 cases.\n\n"
                "DEMO PATH: add a note, watch it go stale, ask the chatbot, delete it; Blocklist decline and override; Rules tab backtest.")
cy, ch = 1.42, 5.53
tag(s, 0.5, cy, 1.5, "BUILT IN PROTOTYPE")
box(s, 0.5, cy + 0.3, 4.1, ch - 0.3, [head("How an investigator trusts, checks and overrides", 12.5),
    {"runs": [("AI can: ", True, BRAND), ("rank and explain, recommend an action, draft an escalation, answer questions using tools.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("Only a human can: ", True, RED), ("clear, escalate, mark fraud, delete a note. Supervisors alone: override a decline, activate a hard rule.", False, INK)], "size": 10.5, "after": 7},
    bl("Rule score and AI score side by side", 10.5, 4),
    bl("Every indicator cites field = value; accept ✓ or reject ✗, and a rejected finding is not re-asserted", 10.5, 4),
    bl("Full agent trace and guardrail removals visible", 10.5, 4),
    bl("A note changes the next assessment; deleting it purges it everywhere", 10.5, 4),
    bl("Every action needs a reason and lands in the audit log", 10.5, 4),
    bl("Bulk clear previews what is blocked, and why", 10.5, 4),
    bl("“Mark as FRAUD” teaches the system: lookalikes get flagged (R17)", 10.5, 0)])
tag(s, 4.75, cy, 1.5, "BUILT IN PROTOTYPE")
box(s, 4.75, cy + 0.3, 4.1, ch - 0.3, [head("Trade-offs and risks I see", 12.5),
    {"runs": [("Accuracy. ", True, INK), ("No labels, so no accuracy claim. Backtests use synthetic labels (leave-one-out) and measure agreement, not fraud.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("False negatives. ", True, RED), ("The costliest risk; only discovered misses can be measured. Guardrails stop the AI clearing anything risky, but do not find what the rules missed.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("Cost and latency. ", True, INK), ("Fewer tokens, more reference dollars ($0.44 vs $0.12). About 20 s per call; free models rate-limit, so some cases show “LLM unavailable”.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("Trust. ", True, INK), ("Cited fields are verified; the prose is not. It can still be subtly wrong.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("Abuse. ", True, INK), ("Notes are untrusted (injection tests pass); bulk clear is guarded; blocklist entries covering over 25% of the queue need confirmation.", False, INK)], "size": 10.5, "after": 5},
    {"runs": [("Fairness. ", True, INK), ("Distance and weekend signals can proxy for rural or shift-work providers.", False, INK)], "size": 10.5, "after": 0}],
    fill=C(0xF4, 0xF5, 0xF8), line=LINE)
tag(s, 9.0, cy, 2.2, "PRODUCTION PATH: NOT BUILT", built=False)
box(s, 9.0, cy + 0.3, 3.83, ch - 0.3, [head("With more time, in priority order", 12.5),
    {"runs": [("1  False-negative safety net. ", True, INK), ("Miss-rate tracking from confirmed outcomes; a look-back that reopens cleared lookalikes; an independent adversarial reviewer; audit sampling with a statistical bound.", False, INK)], "size": 10.5, "after": 6},
    {"runs": [("2  Learn from outcomes. ", True, INK), ("Re-fit weights and thresholds from confirmed fraud.", False, INK)], "size": 10.5, "after": 6},
    {"runs": [("3  Provider / member graph ", True, INK), ("and a labelled backtest, once real IDs and labels exist.", False, INK)], "size": 10.5, "after": 6},
    {"runs": [("4  Four-eyes ", True, INK), ("on large bulk clears; streaming UI.", False, INK)], "size": 10.5, "after": 10},
    {"t": "Prototyped, then deferred: a statistical anomaly detector. On 50 cases its flagged set changed with the random seed, so I would not ship it.", "size": 10, "color": MUT, "after": 0}],
    fill=WHITE, line=MUT, dashed=True)

prs.save(OUT)


# ------------------------------------------------------------------ automated fit check (Arial metrics)
_F = {}


def _font(size, bold):
    key = (round(size * 10), bold)
    if key not in _F:
        _F[key] = ImageFont.truetype(f"/System/Library/Fonts/Supplemental/Arial{' Bold' if bold else ''}.ttf", key[0])
    return _F[key]


def _lines(txt, size, bold, width_pt):
    f, words, n, cur = _font(size, bold), txt.split(" "), 1, ""
    for w in words:
        trial = (cur + " " + w).strip()
        if f.getlength(trial) / 10 <= width_pt or not cur:
            cur = trial
        else:
            n, cur = n + 1, w
    return n


def check(path):
    prs2, bad = Presentation(path), []
    for si, slide in enumerate(prs2.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            x, y, w, h = (sh.left / 914400, sh.top / 914400, sh.width / 914400, sh.height / 914400)
            if x < 0.45 or y < 0.25 or x + w > 13.333 - 0.45 or y + h > 7.5 - 0.1:
                bad.append(f"slide {si}: shape at ({x:.2f},{y:.2f},{w:.2f},{h:.2f}) is outside the margins")
            tf, total = sh.text_frame, (tf_margin := (sh.text_frame.margin_top + sh.text_frame.margin_bottom) / 12700)
            avail = w * 72 - (tf.margin_left + tf.margin_right) / 12700
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                size = p.runs[0].font.size.pt
                bold = any(r.font.bold for r in p.runs)
                marl = int(p._p.pPr.get("marL", 0)) / 12700 if p._p.pPr is not None else 0
                n = _lines("".join(r.text for r in p.runs), size, bold, avail - marl)
                total += n * size * 1.2 + (p.space_after.pt if p.space_after is not None else 0)
            if total > h * 72 * 0.97:
                bad.append(f"slide {si}: text needs {total / 72:.2f} in but shape is {h:.2f} in: \"{tf.text[:50]}...\"")
    return bad


problems = check(OUT)
print("saved", OUT)
print("FIT CHECK:", "OK — all text fits, all shapes inside margins" if not problems else "")
for p in problems:
    print("  PROBLEM", p)
sys.exit(1 if problems else 0)
