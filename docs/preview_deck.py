"""Renders Writeup.pptx to docs/deck_preview.html for visual QA (no LibreOffice on this machine).

    python docs/preview_deck.py

Reads the saved deck's real shape geometry, fills, and runs, and lays them out as absolutely
positioned HTML at 96 px/in. It is an approximation of PowerPoint's renderer — good enough to
catch collisions, crowding and bad rhythm, not a substitute for opening the file.
"""
import base64
import html
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = Path(__file__).resolve().parent
PX = 96
EMU = 914400


def px(v):
    return v / EMU * PX


def rgb(c):
    try:
        return f"#{c.rgb}"
    except Exception:
        return "transparent"


def run_html(r):
    st = [f"font-size:{r.font.size.pt}pt" if r.font.size else "", "font-weight:700" if r.font.bold else "",
          f"color:{rgb(r.font.color)}" if r.font.color and r.font.color.type is not None else ""]
    return f"<span style=\"{';'.join(s for s in st if s)}\">{html.escape(r.text)}</span>"


ALIGN = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}
ANCHOR = {MSO_ANCHOR.MIDDLE: "center", MSO_ANCHOR.BOTTOM: "flex-end"}

out = ["<!doctype html><meta charset='utf-8'><title>Deck preview</title>",
       "<style>body{background:#334155;margin:0;padding:24px;font-family:Arial,Helvetica,sans-serif}",
       ".slide{position:relative;width:%dpx;height:%dpx;background:#fff;margin:0 auto 26px;overflow:hidden;"
       "box-shadow:0 8px 30px rgba(0,0,0,.35)}" % (13.333 * PX, 7.5 * PX),
       ".sh{position:absolute;box-sizing:border-box;display:flex;flex-direction:column}",
       "p{margin:0;line-height:1.2}</style>"]

prs = Presentation(HERE / "Writeup.pptx")
for slide in prs.slides:
    out.append("<div class='slide'>")
    for sh in slide.shapes:
        style = [f"left:{px(sh.left):.1f}px", f"top:{px(sh.top):.1f}px",
                 f"width:{px(sh.width):.1f}px", f"height:{px(sh.height):.1f}px"]
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            b64 = base64.b64encode(sh.image.blob).decode()
            out.append(f"<img class='sh' style=\"{';'.join(style)}\" src='data:image/png;base64,{b64}'>")
            continue
        try:
            if sh.fill.type is not None and sh.fill.type == 1:
                style.append(f"background:{rgb(sh.fill.fore_color)}")
        except Exception:
            pass
        try:
            if sh.line.fill.type == 1:
                style.append(f"border:1px solid {rgb(sh.line.color)}")
        except Exception:
            pass
        if sh.shape_type is not None and "ROUNDED" in str(sh.shape_type):
            style.append("border-radius:7px")
        body = ""
        if sh.has_text_frame:
            tf = sh.text_frame
            style.append(f"padding:{px(tf.margin_top):.0f}px {px(tf.margin_right):.0f}px "
                         f"{px(tf.margin_bottom):.0f}px {px(tf.margin_left):.0f}px")
            style.append(f"justify-content:{ANCHOR.get(tf.vertical_anchor, 'flex-start')}")
            ps = []
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                pst = [f"text-align:{ALIGN.get(p.alignment, 'left')}",
                       f"margin-bottom:{p.space_after.pt if p.space_after else 0}pt",
                       f"line-height:{(p.line_spacing or 1.0) * 1.2}"]
                bullet = "– " if p._p.pPr is not None and p._p.pPr.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None else ""
                ps.append(f"<p style=\"{';'.join(pst)}\">{bullet}{''.join(run_html(r) for r in p.runs)}</p>")
            body = "".join(ps)
        out.append(f"<div class='sh' style=\"{';'.join(style)}\">{body}</div>")
    out.append("</div>")

dst = HERE / "deck_preview.html"
dst.write_text("\n".join(out), encoding="utf-8")
print("wrote", dst)
